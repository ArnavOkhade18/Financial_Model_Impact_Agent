from openpyxl import load_workbook

def extract_excel_data(excel_file):

    wb = load_workbook(excel_file, data_only=False)

    output = ""

    for sheet in wb.sheetnames:

        ws = wb[sheet]

        output += f"\n\n===== SHEET: {sheet} =====\n"

        for row in ws.iter_rows():

            for cell in row:

                if cell.value is not None:

                    output += (
                        f"{cell.coordinate} | "
                        f"{cell.value}\n"
                    )

    return output